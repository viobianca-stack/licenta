# -*- coding: utf-8 -*-
"""
Concise license defense presentation (10 slides, ~7 min + demo).
Styled after the 'AI Technology Project Proposal' Slidesgo template:
dark purple background (#27173A), cyan/purple/blue accents, decorative blobs.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import lxml.etree as etree

# ─── Template palette ─────────────────────────────────────────────────────────
BG_DARK    = RGBColor(0x27, 0x17, 0x3A)   # dark purple background
BG_CARD    = RGBColor(0x35, 0x21, 0x52)   # lighter purple card
BG_CARD2   = RGBColor(0x3D, 0x26, 0x5E)   # card hover
CYAN       = RGBColor(0x76, 0xF3, 0xFB)   # accent cyan
PURPLE     = RGBColor(0xAE, 0x77, 0xD6)   # accent purple
BLUE       = RGBColor(0x72, 0x9B, 0xCB)   # accent blue
DEEP_PUR   = RGBColor(0x5A, 0x13, 0x87)   # deep purple
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT      = RGBColor(0xE6, 0xDC, 0xF5)   # light lavender text
GREY       = RGBColor(0xA8, 0x96, 0xC4)   # muted lavender
GREEN      = RGBColor(0x5D, 0xF0, 0xB0)   # success green
ORANGE     = RGBColor(0xFF, 0xB4, 0x6B)   # warning

ASSETS = "/home/user/licenta/.assets"
W, H = Inches(10), Inches(5.625)
FONT = "Calibri"

PRES_TITLE = "Detectarea mesajelor generate de AI pe rețelele sociale"
TOTAL = 10


def new_prs():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color=BG_DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def pic(slide, path, x, y, w=None, h=None):
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def rect(slide, x, y, w, h, fill=None, line=None, lw=0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def txt(slide, x, y, w, h, text, size=14, bold=False, italic=False,
        color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True,
        spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = FONT; r.font.color.rgb = color
    return tb


def bullets(slide, x, y, w, h, items, size=12, color=LIGHT, gap=6,
            bullet_color=CYAN, spacing=1.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = Pt(2)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        rb = p.add_run(); rb.text = "● "
        rb.font.size = Pt(size-2); rb.font.color.rgb = bullet_color; rb.font.name = FONT
        rt = p.add_run(); rt.text = item
        rt.font.size = Pt(size); rt.font.color.rgb = color; rt.font.name = FONT
    return tb


def header(slide, kicker, title):
    """Top header: small cyan kicker + white title + underline accent."""
    txt(slide, 0.55, 0.30, 9, 0.3, kicker.upper(), size=12, bold=True,
        color=CYAN, align=PP_ALIGN.LEFT)
    txt(slide, 0.55, 0.55, 9, 0.55, title, size=23, bold=True, color=WHITE)
    rect(slide, 0.57, 1.13, 0.85, 0.045, fill=CYAN)


def footer(slide, n):
    txt(slide, 0.55, 5.32, 6, 0.25, PRES_TITLE, size=8, italic=True, color=GREY)
    txt(slide, 9.0, 5.32, 0.8, 0.25, f"{n}/{TOTAL}", size=9, color=GREY,
        align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, fill=BG_CARD, accent=None):
    r = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if accent:
        rect(slide, x, y, 0.07, h, fill=accent)
    return r


# ─── SLIDES ───────────────────────────────────────────────────────────────────

def s1_title(prs):
    slide = add_slide(prs); set_bg(slide)
    # Decorative blob top-right + particle cloud bottom-left
    pic(slide, f"{ASSETS}/image3.png", 5.1, -1.4, w=6.5)
    pic(slide, f"{ASSETS}/image4.png", -2.2, 2.6, w=5.5)

    txt(slide, 0.7, 1.05, 8.6, 0.35,
        "UNIVERSITATEA POLITEHNICA BUCUREȘTI · FACULTATEA AUTOMATICĂ ȘI CALCULATOARE",
        size=10, bold=True, color=GREY)
    txt(slide, 0.7, 1.45, 8.6, 0.35, "LUCRARE DE DIPLOMĂ",
        size=13, bold=True, color=CYAN)
    txt(slide, 0.7, 1.95, 8.6, 1.5, PRES_TITLE,
        size=34, bold=True, color=WHITE, spacing=1.0)

    rect(slide, 0.73, 3.55, 0.7, 0.05, fill=PURPLE)
    txt(slide, 0.7, 3.75, 5, 0.3, "Absolvent", size=10, color=GREY)
    txt(slide, 0.7, 4.0, 5, 0.35, "Bianca-Ștefania GHEORGHE",
        size=15, bold=True, color=WHITE)
    txt(slide, 0.7, 4.5, 6, 0.3, "Coordonator", size=10, color=GREY)
    txt(slide, 0.7, 4.75, 6.5, 0.35, "Conf.dr.ing. Daniel-Marian MEREZEANU",
        size=15, bold=True, color=WHITE)
    txt(slide, 7.4, 4.9, 2.3, 0.3, "Sesiunea Iunie 2026",
        size=11, color=CYAN, align=PP_ALIGN.RIGHT)


def s2_tema(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image5.png", 7.6, -1.0, w=3.3)
    header(slide, "Tema lucrării", "Context și stadiul actual")
    footer(slide, 2)

    # Left: context
    card(slide, 0.55, 1.35, 4.5, 3.75, accent=PURPLE)
    txt(slide, 0.75, 1.5, 4.1, 0.35, "Problema", size=14, bold=True, color=PURPLE)
    bullets(slide, 0.75, 1.95, 4.15, 3.1, [
        "Modelele LLM (ChatGPT, GPT-4, Gemini) generează texte aproape indistinguibile de cele scrise de oameni.",
        "Rețelele sociale: publicare instantanee, fără filtrare editorială → dezinformarea se propagă rapid.",
        "Texte scurte și informale (abrevieri, argou, greșeli intenționate) — metodele clasice funcționează slab aici.",
        "Impact real: site-uri de știri false AI ×10 în 2023 (49→600+); 64% se tem de manipularea alegerilor (KPMG 2025).",
    ], size=12)

    # Right: state of the art
    card(slide, 5.25, 1.35, 4.2, 3.75, accent=CYAN)
    txt(slide, 5.45, 1.5, 3.8, 0.35, "Stadiul actual al cercetării (din 2019)",
        size=14, bold=True, color=CYAN)
    bullets(slide, 5.45, 1.95, 3.85, 2.0, [
        "Metode statistice — perplexitate, entropie (GLTR)",
        "Watermarking — semnal ascuns la generare",
        "Clasificatori supervizați — RoBERTa fine-tuned",
        "Metode zero-shot — DetectGPT",
    ], size=12, bullet_color=PURPLE)
    rect(slide, 5.45, 4.05, 3.8, 0.02, fill=GREY)
    txt(slide, 5.45, 4.15, 3.85, 0.95,
        "Limitare comună: toate sunt evaluate pe texte lungi și formale, iar parafrazarea reduce acuratețea de la >90% la <50%.",
        size=11.5, italic=True, color=ORANGE, spacing=1.05)


def s3_obiective(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image4.png", 6.8, 1.3, w=4.0)
    header(slide, "Obiective", "Ce mi-am propus să realizez")
    footer(slide, 3)

    goals = [
        (PURPLE, "Analiză critică", "Studierea metodelor existente de detectare și a limitărilor lor pe rețelele sociale."),
        (CYAN, "Soluție specializată", "Detector bazat pe machine learning, antrenat specific pe texte scurte de social media (50-500 caractere)."),
        (BLUE, "Comparație de modele", "Evaluarea sistematică a 3 clasificatori complementari pe același set de date."),
        (GREEN, "Aplicație accesibilă", "Interfață web cu justificarea transparentă a deciziilor, nu doar un scor final."),
    ]
    for i, (c, t, d) in enumerate(goals):
        y = 1.45 + i * 0.92
        card(slide, 0.55, y, 5.9, 0.82, accent=c)
        rect(slide, 0.78, y+0.21, 0.42, 0.42, fill=c, shape=MSO_SHAPE.OVAL)
        txt(slide, 0.78, y+0.22, 0.42, 0.4, str(i+1), size=18, bold=True,
            color=BG_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, 1.4, y+0.10, 4.9, 0.32, t, size=14, bold=True, color=c)
        txt(slide, 1.4, y+0.42, 4.95, 0.38, d, size=10.5, color=LIGHT, spacing=1.0)


def s4_tehnologii(prs):
    slide = add_slide(prs); set_bg(slide)
    header(slide, "Implementare", "Tehnologii utilizate")
    footer(slide, 4)

    cols = [
        (PURPLE, "Date", ["Python 3.10", "pandas — filtrare/echilibrare", "datasets (HuggingFace) — streaming"]),
        (CYAN, "ML clasic", ["scikit-learn — NB, Reg. Logistică", "TF-IDF (10.000, n-grame 1-2)", "NLTK — cuvinte stop"]),
        (BLUE, "Deep Learning / NLP", ["PyTorch 2.x — GPU", "Transformers — RoBERTa-base", "SpaCy — NER"]),
        (GREEN, "Aplicație web", ["Flask 3.x — API REST", "HTML / CSS / JS", "Chart.js — vizualizări"]),
    ]
    cw = 2.18
    for i, (c, t, items) in enumerate(cols):
        x = 0.55 + i * (cw + 0.12)
        card(slide, x, 1.4, cw, 2.55)
        rect(slide, x, 1.4, cw, 0.5, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, x, 1.65, cw, 0.25, fill=c)
        txt(slide, x, 1.45, cw, 0.4, t, size=13, bold=True, color=BG_DARK,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        bullets(slide, x+0.15, 2.05, cw-0.25, 1.85, items, size=10.5,
                bullet_color=c, gap=5)

    # Bottom hardware bar
    card(slide, 0.55, 4.15, 8.9, 0.95, fill=BG_CARD2, accent=ORANGE)
    txt(slide, 0.78, 4.25, 4, 0.3, "Mediu de antrenare", size=12, bold=True, color=ORANGE)
    txt(slide, 0.78, 4.55, 8.5, 0.5,
        "Google Colaboratory · GPU NVIDIA Tesla T4 (15.6 GB VRAM) · precizie mixtă FP16 · "
        "antrenare RoBERTa ~85 min (3 epoci) · modele clasice < 2 min · inferență locală pe CPU (0.5-2 s/text).",
        size=10.5, color=LIGHT, spacing=1.05)


def s5_analiza(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image3.png", 7.0, -1.3, w=4.0)
    header(slide, "Analiză preliminară", "Ce am analizat pentru a stabili specificațiile")
    footer(slide, 5)

    items = [
        (PURPLE, "Instrumente existente", "5 aplicații comerciale studiate (GPTZero, ZeroGPT, OpenAI Classifier, Originality.ai, Sapling) — toate slabe pe texte scurte."),
        (CYAN, "Metode din literatură", "4 categorii identificate (statistice, watermarking, supervizate, zero-shot) și limitările lor."),
        (BLUE, "Specificul social media", "Lungime mică, limbaj informal, abrevieri și greșeli intenționate îngreunează clasificarea."),
        (GREEN, "Concluzii → specificații", "Set de date 50-500 caractere · comparație multi-model · justificare per cuvânt · procesare în lot."),
    ]
    for i, (c, t, d) in enumerate(items):
        x = 0.55 + (i % 2) * 4.55
        y = 1.45 + (i // 2) * 1.78
        card(slide, x, y, 4.35, 1.6, accent=c)
        txt(slide, x+0.25, y+0.18, 3.95, 0.35, t, size=14, bold=True, color=c)
        rect(slide, x+0.25, y+0.58, 3.9, 0.02, fill=GREY)
        txt(slide, x+0.25, y+0.68, 3.95, 0.85, d, size=11.5, color=LIGHT, spacing=1.1)


def s6_proiectare(prs):
    slide = add_slide(prs); set_bg(slide)
    header(slide, "Proiectare", "Decizii de arhitectură și model de date")
    footer(slide, 6)

    # Architecture (client-server) left
    card(slide, 0.55, 1.35, 4.5, 3.75, accent=PURPLE)
    txt(slide, 0.75, 1.5, 4.1, 0.35, "Arhitectură software: client-server",
        size=13, bold=True, color=PURPLE)
    bullets(slide, 0.75, 1.95, 4.15, 1.5, [
        "Backend Flask: încarcă cele 3 modele la pornire; endpoint-uri /predict și /batch.",
        "Frontend HTML/CSS/JS: vizualizări, comutare model în browser fără re-apel server.",
    ], size=11.5)
    txt(slide, 0.75, 3.5, 4.1, 0.3, "Trei modele complementare", size=13, bold=True, color=CYAN)
    bullets(slide, 0.75, 3.9, 4.15, 1.15, [
        "Naive Bayes & Reg. Logistică (TF-IDF) — baseline rapid, interpretabil.",
        "RoBERTa fine-tuned (125M parametri) — model principal contextual.",
    ], size=11.5, bullet_color=PURPLE)

    # Data model right
    card(slide, 5.25, 1.35, 4.2, 3.75, accent=BLUE)
    txt(slide, 5.45, 1.5, 3.8, 0.35, "Model de date: 6 surse combinate",
        size=13, bold=True, color=BLUE)
    rows = [
        ("HC3 Reddit ELI5", "perechi uman/AI"),
        ("TweetEval", "tweets reale"),
        ("Reddit comentarii", "limbaj informal"),
        ("RAID (2024)", "AI multi-model"),
        ("AI_Human.csv", "volum/diversitate"),
        ("AI Detection", "texte AI extra"),
    ]
    for i, (src, desc) in enumerate(rows):
        y = 1.98 + i * 0.36
        txt(slide, 5.45, y, 2.1, 0.32, "● " + src, size=11, color=WHITE, bold=True)
        txt(slide, 7.5, y, 1.85, 0.32, desc, size=10, color=GREY, align=PP_ALIGN.RIGHT)
    rect(slide, 5.45, 4.18, 3.8, 0.02, fill=GREY)
    txt(slide, 5.45, 4.28, 3.85, 0.8,
        "45.834 exemple echilibrate 50/50 (uman/AI) · filtrate 50-500 caractere · "
        "împărțire stratificată 70/15/15.",
        size=11, color=GREEN, spacing=1.1)


def s7_realizare(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image4.png", 7.2, 2.7, w=3.6)
    header(slide, "Realizare practică", "Rezultate obținute  →  urmează demo")
    footer(slide, 7)

    # Results table (3 models)
    models = [
        ("Naive Bayes", "94.24%", "94.17%", PURPLE),
        ("Reg. Logistică", "96.71%", "96.66%", BLUE),
        ("RoBERTa", "99.17%", "99.18%", GREEN),
    ]
    # header row
    tx = 0.55
    rect(slide, tx, 1.4, 5.4, 0.42, fill=DEEP_PUR, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(slide, tx+0.2, 1.43, 2.0, 0.36, "Model", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    txt(slide, tx+2.7, 1.43, 1.3, 0.36, "Acuratețe", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(slide, tx+3.9, 1.43, 1.3, 0.36, "F1-score", size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    for i, (name, acc, f1, c) in enumerate(models):
        y = 1.88 + i * 0.52
        best = (i == 2)
        rect(slide, tx, y, 5.4, 0.46, fill=BG_CARD2 if best else BG_CARD,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, tx, y, 0.06, 0.46, fill=c)
        txt(slide, tx+0.2, y+0.04, 2.4, 0.4, name, size=12, bold=best, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, tx+2.7, y+0.04, 1.3, 0.4, acc, size=13 if best else 12, bold=best,
            color=c if best else LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, tx+3.9, y+0.04, 1.3, 0.4, f1, size=13 if best else 12, bold=best,
            color=c if best else LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    txt(slide, 0.55, 3.55, 5.4, 0.6,
        "Testat pe 6.876 exemple (50/50). Pe sursele propriu-zise de social media: acuratețe >99.5%.",
        size=11, italic=True, color=GREY, spacing=1.05)

    # Demo card right
    card(slide, 6.25, 1.4, 3.2, 2.4, accent=CYAN)
    txt(slide, 6.45, 1.52, 2.9, 0.3, "Demo aplicație", size=13, bold=True, color=CYAN)
    bullets(slide, 6.45, 1.92, 2.85, 1.85, [
        "Verdict + scor agregat de risc (0-100)",
        "Evidențiere cuvinte (Leave-One-Out)",
        "Comparație 3 modele simultan",
        "Procesare în lot CSV",
    ], size=10.5, bullet_color=GREEN, gap=5)

    # Confusion matrix mini
    card(slide, 0.55, 4.25, 8.9, 0.85, fill=BG_CARD2, accent=GREEN)
    txt(slide, 0.78, 4.34, 8.5, 0.7,
        "Confusion matrix RoBERTa: 3.388 uman corect · 3.431 AI corect · 50 fals-pozitive · 7 fals-negative — "
        "aproape toate erorile provin din texte formale, nu din social media.",
        size=10.5, color=LIGHT, spacing=1.05)


def s8_contributii(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image5.png", 7.8, 3.2, w=2.6)
    header(slide, "Sinteză", "Contribuții originale")
    footer(slide, 8)

    contribs = [
        (PURPLE, "Set de date dedicat", "6 surse de social media combinate, focalizat pe texte scurte (50-500 car.)."),
        (CYAN, "Comparație 3 modele", "Naive Bayes, Reg. Logistică și RoBERTa evaluate în condiții identice."),
        (BLUE, "Justificare Leave-One-Out", "Explicare la nivel de cuvânt — absentă în instrumentele comerciale."),
        (GREEN, "Scor agregat de risc", "Indicator 0-100 multi-semnal (RoBERTa 60% + fraze AI + statistici vocabular)."),
        (ORANGE, "Aplicație web integrată", "Analiză, comparație, per-propoziție, NER, batch CSV — open-source."),
    ]
    for i, (c, t, d) in enumerate(contribs):
        y = 1.4 + i * 0.74
        rect(slide, 0.55, y, 0.5, 0.62, fill=c, shape=MSO_SHAPE.OVAL)
        txt(slide, 0.55, y+0.13, 0.5, 0.4, str(i+1), size=18, bold=True,
            color=BG_DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        card(slide, 1.2, y, 7.6, 0.62)
        txt(slide, 1.4, y+0.06, 3.0, 0.5, t, size=13, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
        txt(slide, 4.3, y+0.06, 4.35, 0.5, d, size=10.5, color=LIGHT,
            anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)


def s9_dezvoltari(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image3.png", 6.9, -1.4, w=4.0)
    header(slide, "Perspective", "Dezvoltări viitoare")
    footer(slide, 9)

    items = [
        (PURPLE, "Extindere multilingvă", "XLM-RoBERTa sau seturi de date dedicate, inclusiv pentru limba română."),
        (CYAN, "Reantrenare periodică", "Pe modele noi (GPT-4o, Claude, Gemini) — combaterea „concept drift”."),
        (BLUE, "Robustețe la evitare", "Testare sistematică împotriva parafrazării și „humanizării” textului."),
        (GREEN, "Semnale comportamentale", "Integrarea metadatelor de rețea (frecvență postare, profil cont)."),
        (ORANGE, "Infrastructură", "Extensie de browser, API public REST, containerizare Docker."),
        (PURPLE, "Calibrare scor", "Optimizarea automată a ponderilor scorului agregat de risc."),
    ]
    for i, (c, t, d) in enumerate(items):
        x = 0.55 + (i % 2) * 4.55
        y = 1.45 + (i // 2) * 1.18
        card(slide, x, y, 4.35, 1.02, accent=c)
        txt(slide, x+0.25, y+0.13, 3.95, 0.32, t, size=13, bold=True, color=c)
        txt(slide, x+0.25, y+0.5, 3.95, 0.45, d, size=10.5, color=LIGHT, spacing=1.0)


def s10_final(prs):
    slide = add_slide(prs); set_bg(slide)
    pic(slide, f"{ASSETS}/image4.png", -1.8, -1.5, w=5.0)
    pic(slide, f"{ASSETS}/image3.png", 6.0, 2.8, w=5.0)

    txt(slide, 0.7, 1.6, 8.6, 1.0, "Vă mulțumesc pentru atenție!",
        size=38, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(slide, 0.7, 2.75, 8.6, 0.5, "Întrebări?",
        size=22, color=CYAN, align=PP_ALIGN.CENTER)

    stats = [("99.17%", "Acuratețe"), ("45.834", "Exemple"), ("3", "Modele"), ("5", "Contribuții")]
    for i, (v, l) in enumerate(stats):
        x = 1.55 + i * 1.78
        card(slide, x, 3.7, 1.55, 1.1, fill=BG_CARD)
        txt(slide, x, 3.82, 1.55, 0.5, v, size=22, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
        txt(slide, x, 4.35, 1.55, 0.35, l, size=10, color=LIGHT, align=PP_ALIGN.CENTER)


def build():
    prs = new_prs()
    s1_title(prs); s2_tema(prs); s3_obiective(prs); s4_tehnologii(prs)
    s5_analiza(prs); s6_proiectare(prs); s7_realizare(prs); s8_contributii(prs)
    s9_dezvoltari(prs); s10_final(prs)
    out = "/home/user/licenta/Prezentare_Licenta_Gheorghe_Bianca.pptx"
    prs.save(out)
    print("Saved:", out, "| slides:", len(prs.slides))


if __name__ == "__main__":
    build()
